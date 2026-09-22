"""
report_generator.py
---------------------
Takes the CLEANED data + cleaning log and automatically builds:
    1. An Excel report (output/Sales_Report.xlsx) with:
        - Summary sheet (KPIs + cleaning log)
        - Cleaned Data sheet
        - Native Excel charts (bar + pie)
    2. A PNG visual summary dashboard (output/summary_dashboard.png)

Usage:
    from report_generator import generate_report
    generate_report(df_clean, cleaning_log)
"""

import pandas as pd
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

from openpyxl import Workbook
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import BarChart, PieChart, Reference


def _style_header(ws, row_idx, ncols):
    for col in range(1, ncols + 1):
        cell = ws.cell(row=row_idx, column=col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")


def generate_report(df: pd.DataFrame, log: dict,
                     excel_path="output/Sales_Report.xlsx",
                     png_path="output/summary_dashboard.png"):

    # =========================================================
    # PART A: Excel report
    # =========================================================
    wb = Workbook()

    # ---- Sheet 1: Summary ----
    ws1 = wb.active
    ws1.title = "Summary"
    ws1["A1"] = "Data Cleaning & Sales Summary Report"
    ws1["A1"].font = Font(bold=True, size=14)
    ws1.merge_cells("A1:D1")

    ws1["A3"] = "Cleaning Log"
    ws1["A3"].font = Font(bold=True, size=12)
    r = 4
    for key, val in log.items():
        ws1.cell(row=r, column=1, value=key.replace("_", " ").title())
        ws1.cell(row=r, column=2, value=val)
        r += 1

    r += 1
    ws1.cell(row=r, column=1, value="Key Metrics").font = Font(bold=True, size=12)
    r += 1
    total_sales = df["TotalSale"].sum() if "TotalSale" in df.columns else 0
    avg_order = df["TotalSale"].mean() if "TotalSale" in df.columns else 0
    ws1.cell(row=r, column=1, value="Total Sales Value")
    ws1.cell(row=r, column=2, value=round(float(total_sales), 2))
    r += 1
    ws1.cell(row=r, column=1, value="Average Order Value")
    ws1.cell(row=r, column=2, value=round(float(avg_order), 2))
    r += 1
    ws1.cell(row=r, column=1, value="Total Orders (cleaned)")
    ws1.cell(row=r, column=2, value=len(df))

    for col, width in zip("ABCD", (30, 18, 14, 14)):
        ws1.column_dimensions[col].width = width

    # ---- Sheet 2: Cleaned Data ----
    ws2 = wb.create_sheet("Cleaned Data")
    for row in dataframe_to_rows(df, index=False, header=True):
        ws2.append(row)
    _style_header(ws2, 1, len(df.columns))
    for i, col in enumerate(df.columns, start=1):
        ws2.column_dimensions[chr(64 + i) if i <= 26 else "A"].width = 16

    # ---- Sheet 3: Region-wise summary + chart ----
    ws3 = wb.create_sheet("Region Summary")
    if "Region" in df.columns and "TotalSale" in df.columns:
        region_summary = (
            df.groupby("Region", as_index=False)["TotalSale"]
            .sum()
            .sort_values("TotalSale", ascending=False)
        )
        ws3.append(["Region", "TotalSale"])
        _style_header(ws3, 1, 2)
        for _, row in region_summary.iterrows():
            ws3.append([row["Region"], round(float(row["TotalSale"]), 2)])

        chart = BarChart()
        chart.title = "Total Sales by Region"
        chart.x_axis.title = "Region"
        chart.y_axis.title = "Total Sales"
        data_ref = Reference(ws3, min_col=2, min_row=1, max_row=len(region_summary) + 1)
        cats_ref = Reference(ws3, min_col=1, min_row=2, max_row=len(region_summary) + 1)
        chart.add_data(data_ref, titles_from_data=True)
        chart.set_categories(cats_ref)
        ws3.add_chart(chart, "E2")

    # ---- Sheet 4: Product summary + pie chart ----
    ws4 = wb.create_sheet("Product Summary")
    if "Product" in df.columns and "TotalSale" in df.columns:
        product_summary = (
            df.groupby("Product", as_index=False)["TotalSale"]
            .sum()
            .sort_values("TotalSale", ascending=False)
        )
        ws4.append(["Product", "TotalSale"])
        _style_header(ws4, 1, 2)
        for _, row in product_summary.iterrows():
            ws4.append([row["Product"], round(float(row["TotalSale"]), 2)])

        pie = PieChart()
        pie.title = "Sales Share by Product"
        data_ref = Reference(ws4, min_col=2, min_row=1, max_row=len(product_summary) + 1)
        cats_ref = Reference(ws4, min_col=1, min_row=2, max_row=len(product_summary) + 1)
        pie.add_data(data_ref, titles_from_data=True)
        pie.set_categories(cats_ref)
        ws4.add_chart(pie, "E2")

    wb.save(excel_path)

    # =========================================================
    # PART B: PNG dashboard (quick visual summary)
    # =========================================================
    fig, axes = plt.subplots(2, 2, figsize=(12, 8))
    fig.suptitle("Sales Data - Visual Summary", fontsize=16, fontweight="bold")

    if "Region" in df.columns and "TotalSale" in df.columns:
        region_summary.plot(kind="bar", x="Region", y="TotalSale", ax=axes[0, 0], legend=False, color="#4472C4")
        axes[0, 0].set_title("Total Sales by Region")
        axes[0, 0].set_ylabel("Sales")

    if "Product" in df.columns and "TotalSale" in df.columns:
        axes[0, 1].pie(product_summary["TotalSale"], labels=product_summary["Product"], autopct="%1.1f%%")
        axes[0, 1].set_title("Sales Share by Product")

    if "OrderDate" in df.columns and "TotalSale" in df.columns:
        monthly = df.set_index("OrderDate").resample("ME")["TotalSale"].sum()
        monthly.plot(ax=axes[1, 0], marker="o", color="#ED7D31")
        axes[1, 0].set_title("Monthly Sales Trend")
        axes[1, 0].set_ylabel("Sales")

    if "SalesRep" in df.columns and "TotalSale" in df.columns:
        rep_summary = df.groupby("SalesRep")["TotalSale"].sum().sort_values(ascending=False)
        rep_summary.plot(kind="bar", ax=axes[1, 1], color="#70AD47")
        axes[1, 1].set_title("Sales by Rep")
        axes[1, 1].set_ylabel("Sales")

    plt.tight_layout(rect=[0, 0, 1, 0.95])
    plt.savefig(png_path, dpi=150)
    plt.close(fig)

    print(f"Excel report saved to: {excel_path}")
    print(f"Visual dashboard saved to: {png_path}")


if __name__ == "__main__":
    from data_cleaning import clean_data

    raw = pd.read_csv("data/raw_sales_data.csv")
    cleaned, log = clean_data(raw)
    generate_report(cleaned, log)
